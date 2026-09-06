from __future__ import annotations

from io import BytesIO

import pytest

from app.services.record_pipeline.document_parser import (
    UnsupportedDocumentError,
    extract_text_from_document,
)


def test_extract_text_from_txt_utf8():
    content = "안녕하세요\n반갑습니다".encode("utf-8")
    assert extract_text_from_document(content, "note.txt") == "안녕하세요\n반갑습니다"


def test_extract_text_from_txt_falls_back_to_cp949():
    # Windows Notepad-era Korean text files are often CP949, not UTF-8 — bytes
    # that are invalid UTF-8 must still come back as the original text.
    content = "메모장에서 저장한 텍스트".encode("cp949")
    assert extract_text_from_document(content, "note.txt") == "메모장에서 저장한 텍스트"


def test_extract_text_from_md_is_treated_as_plain_text():
    content = "# 제목\n\n본문 내용입니다.".encode("utf-8")
    assert extract_text_from_document(content, "notes.md") == "# 제목\n\n본문 내용입니다."


def test_extract_text_from_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("첫 번째 문단입니다.")
    doc.add_paragraph("두 번째 문단입니다.")
    buf = BytesIO()
    doc.save(buf)

    text = extract_text_from_document(buf.getvalue(), "resume.docx")
    assert "첫 번째 문단입니다." in text
    assert "두 번째 문단입니다." in text


def test_extract_text_from_malformed_docx_raises_unsupported():
    with pytest.raises(UnsupportedDocumentError):
        extract_text_from_document(b"not a real docx file", "resume.docx")


def test_extract_text_from_malformed_hwp_raises_unsupported():
    with pytest.raises(UnsupportedDocumentError):
        extract_text_from_document(b"not a real hwp file", "resume.hwp")


def test_unsupported_extension_raises():
    with pytest.raises(UnsupportedDocumentError):
        extract_text_from_document(b"whatever", "resume.pdf")
